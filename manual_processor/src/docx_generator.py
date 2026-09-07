"""
Word Document Generator Module
Creates formatted Word documents from processed text
"""

import json
import logging
from pathlib import Path
from typing import Optional, Dict, Any, Tuple

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    Document = None

logger = logging.getLogger(__name__)

DOCX_TEMPLATE_DIR = Path(__file__).parent.parent / "templates" / "docx"


class DocxTemplateLoader:
    """Loads and manages Word document output templates"""

    _cache: Dict[str, Dict[str, Any]] = {}

    @classmethod
    def load_template(cls, name: str) -> Dict[str, Any]:
        """Load a docx template by name"""
        if name in cls._cache:
            return cls._cache[name]

        template_path = DOCX_TEMPLATE_DIR / f"{name}.json"
        if not template_path.exists():
            logger.warning(f"Docx template not found: {name}, using defaults")
            return cls._get_default_template()

        try:
            with open(template_path, "r", encoding="utf-8") as f:
                template = json.load(f)
        except Exception as e:
            logger.warning(f"Failed to load docx template {name}: {e}, using defaults")
            return cls._get_default_template()

        cls._cache[name] = template
        return template

    @classmethod
    def _get_default_template(cls) -> Dict[str, Any]:
        """Return default template settings"""
        return {
            "name": "default",
            "layout": "standard",
            "font_size": 11,
            "compact_mode": False,
            "emoji_enabled": False,
            "title_font_size": 22,
            "header_font_size": 16,
            "subheader_font_size": 14,
            "bullet_indent": 0.25,
            "paragraph_spacing": 6,
            "line_spacing": 1.15,
            "page_break_before_diagram": True
        }

    @classmethod
    def resolve_template(cls, template: Dict[str, Any]) -> Dict[str, Any]:
        """Resolve template inheritance"""
        if not template.get("parent"):
            return template

        parent_name = template["parent"]
        parent_template = cls.load_template(parent_name)
        resolved_parent = cls.resolve_template(parent_template)

        merged = resolved_parent.copy()
        merged.update(template)
        return merged

    @classmethod
    def validate_template(cls, template: Dict[str, Any]) -> Tuple[bool, str]:
        """Validate a template dictionary"""
        valid_layouts = ["standard", "compact", "fancy"]
        if "layout" in template and template["layout"] not in valid_layouts:
            return False, f"Invalid layout: {template['layout']}"
        return True, ""

    @classmethod
    def clear_cache(cls) -> None:
        """Clear the template cache"""
        cls._cache.clear()


class WordGenerator:
    """Generates formatted Word documents"""
    
    def __init__(self):
        """Initialize Word generator"""
        if Document is None:
            raise ImportError("python-docx is required for Word generation. Install with: pip install python-docx")
        logger.debug("WordGenerator initialized")
    
    def insert_diagram_image(self, document: Document, image_path: Path,
                             caption: str = "作業フロー図") -> None:
        """Word文書にフローチャート画像を挿入"""
        if not image_path or not Path(image_path).exists():
            return

        try:
            document.add_heading(f"📊 {caption}", level=1)
            document.add_picture(str(image_path), width=Inches(6.0))
            caption_para = document.add_paragraph("※ 上記フローチャートはAIが自動生成したものです。")
            caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            logger.info(f"Inserted diagram image into Word document: {image_path}")
        except Exception as e:
            logger.warning(f"Word への画像挿入失敗: {e}")

    def generate_word(self, content: str, output_path: Path, title: str = "Processed Manual",
                      compact_layout: bool = False, use_emojis: bool = False,
                      diagram_path: Optional[Path] = None,
                      qr_image_path: Optional[Path] = None,
                      template: Optional[Dict[str, Any]] = None) -> Path:
        """
        Generate a formatted Word document

        Args:
            content: Text content for the document
            output_path: Output file path
            title: Document title
            compact_layout: If True, use compact layout
            use_emojis: If True, insert emojis
            diagram_path: Optional path to diagram image
            qr_image_path: Optional path to QR code image (inserted after title)
            template: Optional template dictionary
        """
        if template:
            t = template
        else:
            t = {
                "font_size": 13 if compact_layout else 11,
                "compact_mode": compact_layout,
                "emoji_enabled": use_emojis,
                "title_font_size": 22,
                "header_font_size": 16,
                "subheader_font_size": 14,
                "paragraph_spacing": 2 if compact_layout else 6,
                "line_spacing": 1.0 if compact_layout else 1.15,
            }

        document = Document()

        use_emojis_flag = t.get("emoji_enabled", use_emojis)

        # Title
        display_title = f"📄 {title}" if use_emojis_flag else title
        heading = document.add_heading(display_title, 0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if t.get("compact_mode"):
            for run in heading.runs:
                run.font.size = Pt(t.get("title_font_size", 22))

        # Insert QR code after title if present
        if qr_image_path and Path(qr_image_path).exists():
            try:
                last_paragraph = document.paragraphs[-1]
                last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = last_paragraph.add_run()
                run.add_picture(str(qr_image_path), width=Inches(1.5))
                qr_caption = document.add_paragraph("🎵 スマホでQRをスキャンして音声再生")
                qr_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                qr_caption.runs[0].font.size = Pt(9)
                qr_caption.runs[0].font.color.rgb = RGBColor(128, 128, 128)
                logger.info(f"Inserted QR image into Word document: {qr_image_path}")
            except Exception as e:
                logger.warning(f"Word へのQR挿入失敗: {e}")

        base_font_size = Pt(t.get("font_size", 11))
        for line in content.split('\n'):
            if line.strip().startswith('## '):
                # Section header
                header_text = line[3:].strip()
                if use_emojis_flag:
                    header_text = f"📌 {header_text}"
                heading = document.add_heading(header_text, level=1)
                if t.get("compact_mode"):
                    for run in heading.runs:
                        run.font.size = Pt(t.get("header_font_size", 16))
                    heading.paragraph_format.space_after = Pt(2)
                    heading.paragraph_format.space_before = Pt(2)
            elif line.strip().startswith('### '):
                # Sub-section header
                header_text = line[4:].strip()
                if use_emojis_flag:
                    header_text = f"📌 {header_text}"
                heading = document.add_heading(header_text, level=2)
                if t.get("compact_mode"):
                    for run in heading.runs:
                        run.font.size = Pt(t.get("subheader_font_size", 14))
                    heading.paragraph_format.space_after = Pt(2)
                    heading.paragraph_format.space_before = Pt(2)
            elif line.strip().startswith('- ') or line.strip().startswith('• '):
                # Bullet point
                bullet_text = line.strip()[2:]
                if use_emojis_flag:
                    bullet_text = f"🔹 {bullet_text}"
                bullet_point = document.add_paragraph(style='List Bullet')
                run = bullet_point.add_run(bullet_text)
                run.font.size = base_font_size
                if t.get("compact_mode"):
                    bullet_point.paragraph_format.space_after = Pt(2)
                    bullet_point.paragraph_format.line_spacing = 1.0
            elif line.strip():
                # Regular paragraph
                paragraph = document.add_paragraph(line.strip())
                if t.get("compact_mode"):
                    for run in paragraph.runs:
                        run.font.size = base_font_size
                    paragraph.paragraph_format.space_after = Pt(t.get("paragraph_spacing", 2))
                    paragraph.paragraph_format.line_spacing = t.get("line_spacing", 1.0)

        # Insert diagram image if present
        if diagram_path and Path(diagram_path).exists():
            self.insert_diagram_image(document, Path(diagram_path))

        # Save document
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))
        logger.info(f"Word document generated: {output_path}")

        return output_path


def create_word_document(content: str, output_path: Path, title: str = "Processed Manual",
                         compact_layout: bool = False, use_emojis: bool = False,
                         diagram_path: Optional[Path] = None,
                         qr_image_path: Optional[Path] = None,
                         template_name: str = None) -> Path:
    """
    Convenience function to generate Word document

    Args:
        content: Text content
        output_path: Output file path
        title: Document title
        compact_layout: If True, use compact layout
        use_emojis: If True, insert emojis
        diagram_path: Optional diagram image path
        qr_image_path: Optional QR code image path
        template_name: Template name to use
    """
    generator = WordGenerator()

    if template_name:
        try:
            template = DocxTemplateLoader.load_template(template_name)
            resolved = DocxTemplateLoader.resolve_template(template)
        except Exception as e:
            logger.warning(f"Docx template loading failed: {e}, using defaults")
            resolved = None
    else:
        resolved = None

    return generator.generate_word(content, output_path, title, compact_layout=compact_layout,
                                   use_emojis=use_emojis, diagram_path=diagram_path,
                                   qr_image_path=qr_image_path,
                                   template=resolved)